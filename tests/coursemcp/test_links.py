"""What goes with what: ``links.py`` and the build side of the mirror.

Pure functions over a directory and slide data -- no git, no app -- plus
``<code>`` and the mirror's paths through the build, slide hyperlinks, a
check against the real hwdesign checkout when one sits beside this repo, and
an opt-in clone from GitHub.  (``plans/unified_materials.md``; the link
graph was first built for demos, ``plans/demo_code_mcp.md``.)
"""

from __future__ import annotations

import os
import subprocess
from pathlib import Path

import pytest

from llmgrader.coursemcp.code import CodeConfig, CodeSync, build_snapshot
from llmgrader.coursemcp.links import (
    OverridesError,
    SlideRef,
    find_references,
    link_items,
    parse_overrides,
    resolve,
)
from test_course_code import CODE_CONFIG, REPO_BINARIES, REPO_FILES

REPO_ROOT = Path(__file__).resolve().parents[2]
HWDESIGN = REPO_ROOT.parent / "hwdesign"

SLIDES = [
    SlideRef("procif", ("Unit 4:  Memory and Processor Interfaces",), 50,
             "See demos/scalar_fun/scalar_fun_vitis"),
    SlideRef("fsm", ("Unit 2:  Sequential Logic and FSMs",), 35, "Run it",
             ("https://example.github.io/hwdesign/docs/demos/stream/",)),
    SlideRef("intro", ("Unit 0:  Introduction",), 6,
             "Course site: https://example.github.io/hwdesign/docs"),
]

# The items the catalog makes from the fixture repo, as their mirror roots.
ROOTS = {
    "procif-demo": ["docs/demos/procif", "demos/scalar_fun"],
    "stream": ["demos/stream", "docs/demos/stream"],
    "fifoif": ["docs/demos/fifoif"],
    "fsm-demo": ["demos/fsm"],
}


def make_tree(root: Path, extra: dict | None = None) -> Path:
    tree = root / "tree"
    for path, text in {**REPO_FILES, **(extra or {})}.items():
        (tree / path).parent.mkdir(parents=True, exist_ok=True)
        (tree / path).write_text(text, encoding="utf-8")
    for path, data in REPO_BINARIES.items():
        (tree / path).parent.mkdir(parents=True, exist_ok=True)
        (tree / path).write_bytes(data)
    return tree


def snapshot_of(tmp_path, extra=None):
    return build_snapshot(make_tree(tmp_path, extra), CodeConfig.from_dict(CODE_CONFIG), "abc123")


def links_of(tmp_path, extra=None, slides=SLIDES):
    return link_items(snapshot_of(tmp_path, extra), slides, ROOTS)


# ---------------------------------------------------------------------------
# References
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("text, expected", [
    ("files are in `hwdesign/demos/stream/avgfilt`.", ("dir", "demos/stream/avgfilt")),
    ("see demos/fsm/counter.sv", ("file", "demos/fsm/counter.sv")),
    ("https://github.com/test/hwdesign/blob/main/demos/fsm/counter.sv", ("file", "demos/fsm/counter.sv")),
    ("https://github.com/test/hwdesign/tree/main/demos/stream/poly", ("dir", "demos/stream/poly")),
    ("https://example.github.io/hwdesign/docs/demos/procif/", ("file", "docs/demos/procif/index.md")),
    ("https://example.github.io/hwdesign/docs/demos/procif/vitis_ip.html",
     ("file", "docs/demos/procif/vitis_ip.md")),
    ("see demos/basic_logic/gone.sv", ("broken", "demos/basic_logic/gone.sv")),
    ("demos/scalar_fun/scalar_fun_vitis/hls_component/impl/ip",
     ("build_output", "demos/scalar_fun/scalar_fun_vitis/hls_component/impl/ip")),
    ("demos/scalar_fun/overlay/design.bit", ("excluded", "demos/scalar_fun/overlay/design.bit")),
    ("https://example.github.io/hwdesign/docs", ("outside", "docs")),
])
def test_each_reference_form_resolves(tmp_path, text, expected) -> None:
    snapshot = snapshot_of(tmp_path)
    [(_, candidates)] = find_references(text, snapshot)
    assert resolve(candidates, snapshot) == expected


def test_a_path_inside_a_longer_path_is_not_a_reference(tmp_path) -> None:
    snapshot = snapshot_of(tmp_path)
    assert find_references("Wrote C:/work/demos/stream/avgfilt/out.csv", snapshot) == []


def test_relative_links_between_docs_pages(tmp_path) -> None:
    snapshot = snapshot_of(tmp_path)
    [(_, candidates)] = find_references("the [example](../procif/) first", snapshot,
                                        page="docs/demos/stream/index.md")
    assert resolve(candidates, snapshot) == ("file", "docs/demos/procif/index.md")


# ---------------------------------------------------------------------------
# Linking items
# ---------------------------------------------------------------------------


def test_slides_link_items_and_give_them_units(tmp_path) -> None:
    links = links_of(tmp_path).items
    assert links["procif-demo"].slides == [{
        "deck": "procif", "slide": 50,
        "via": "slide procif:50 cites demos/scalar_fun/scalar_fun_vitis"}]
    assert links["procif-demo"].units[0]["unit"] == "Unit 4:  Memory and Processor Interfaces"
    # A slide's hyperlink to a docs page links the item that page is in.
    assert [u["unit"] for u in links["stream"].units] == ["Unit 2:  Sequential Logic and FSMs"]
    # A deck named like an item's folder gives the deck's units.
    assert links["fsm-demo"].units[0]["via"] == "deck fsm has the item's name"
    # A site root (the intro slide) names no item.
    assert not any(s["deck"] == "intro" for found in links.values() for s in found.slides)


def test_docs_pages_cite_code_and_relate_items(tmp_path) -> None:
    links = links_of(tmp_path)
    assert links.doc_cites["docs/demos/procif/index.md"] == ["demos/scalar_fun/scalar_fun_vitis"]
    assert links.items["stream"].related == ["procif-demo"]       # ../procif/
    assert links.items["fifoif"].related == ["stream"]            # demos/stream/poly
    assert links.items["procif-demo"].related == []               # its own code


def test_broken_and_build_output_references_are_reported(tmp_path) -> None:
    report = links_of(tmp_path).report
    assert report.broken == [("docs/demos/procif/vitis_ip.md:17", "demos/basic_logic/old_fun.sv")]
    assert "Broken reference" in "\n".join(report.lines())


def test_an_item_from_the_package_is_linked_by_nothing(tmp_path) -> None:
    links = link_items(snapshot_of(tmp_path), SLIDES, {"procif": []})
    assert links.items["procif"].units == [] and links.items["procif"].slides == []


# ---------------------------------------------------------------------------
# The repo's materials.xml
# ---------------------------------------------------------------------------


def test_materials_xml_is_validated() -> None:
    root = parse_overrides("<materials><item id='fsm-demo' unit='Unit 2'>"
                           "<description>x</description><slides deck='fsm'/></item></materials>")
    assert root.find("item").get("id") == "fsm-demo"
    assert parse_overrides("") is None


@pytest.mark.parametrize("text", [
    "<materials><item id='fsm'>",
    "<materials><item id='fsm'><bogus/></item></materials>",
    "<materials><item/></materials>",
    "<materials><item id='fsm' path='demos/fsm'/></materials>",     # it refines, never makes
    "<demos><demo id='fsm'/></demos>",                               # the old file
])
def test_a_broken_materials_xml_is_an_error(text) -> None:
    with pytest.raises(OverridesError):
        parse_overrides(text)


def test_symlinks_are_never_served(tmp_path) -> None:
    tree = make_tree(tmp_path)
    secret = tmp_path / "secret.txt"
    secret.write_text("SECRET", encoding="utf-8")
    try:
        os.symlink(secret, tree / "demos" / "fsm" / "leak.txt")
    except (OSError, NotImplementedError):
        pytest.skip("this system cannot create symlinks")
    snapshot = build_snapshot(tree, CodeConfig.from_dict(CODE_CONFIG), "abc")
    assert "demos/fsm/leak.txt" not in snapshot.files
    assert snapshot.search("SECRET")["hits"] == []


# ---------------------------------------------------------------------------
# The build: <code> and the paths the mirror fetches
# ---------------------------------------------------------------------------


MCP_CONFIG = """<llmgrader_mcp>
  <roots><root id="pub" path="{tree}"/></roots>
  <code repo="{repo}" root="pub">
    <exclude>*.wcfg</exclude>
  </code>
  <materials>
    <kind id="demo" description="Demos"/>
    <kind id="lab" description="Labs"/>
    <item id="procif-demo" kind="demo" source="mirror" docs="docs/demos/procif"
          path="demos/scalar_fun" unit="U"/>
    <items kind="demo" source="mirror" path="demos/*" docs="docs/demos/*"/>
    <items kind="lab" source="mirror" path="labs/*" docs="docs/labs/*"/>
  </materials>
</llmgrader_mcp>
"""


def write_mcp_config(tmp_path, repo="https://github.com/test/hwdesign") -> Path:
    tree = make_tree(tmp_path)
    (tmp_path / "llmgrader_config.xml").write_text(
        "<llmgrader><units><unit><name>U</name><source>u.xml</source></unit></units></llmgrader>",
        encoding="utf-8")
    path = tmp_path / "llmgrader_mcp_config.xml"
    path.write_text(MCP_CONFIG.format(tree=tree.as_posix(), repo=repo), encoding="utf-8")
    return path


def test_the_mirror_fetches_the_paths_the_items_name(tmp_path) -> None:
    from llmgrader.coursemcp.materials import build_materials, load_materials, read_code, read_materials

    config = write_mcp_config(tmp_path)
    spec = read_code(config)
    assert spec.config["repo"] == "https://github.com/test/hwdesign"
    # The explicit item's folders sit inside the patterns' and are dropped.
    assert spec.config["includes"] == [
        {"path": "demos", "format": "text"}, {"path": "docs/demos", "format": "page"},
        {"path": "labs", "format": "text"}, {"path": "docs/labs", "format": "page"}]
    assert spec.config["excludes"] == ["*.wcfg"]
    assert spec.local_root == (tmp_path / "tree").resolve()
    package = tmp_path / "pkg"
    build_materials(read_materials(config), package, code=spec.config, log=lambda *_: None)
    loaded = load_materials(str(package))
    assert loaded.code == spec.config
    assert [i["id"] for i in loaded.mirror_items] == ["procif-demo"]
    assert [p["path"] for p in loaded.patterns] == ["demos/*", "labs/*"]


def test_sparse_patterns_match_a_file_or_a_folder() -> None:
    config = CodeConfig.from_dict({**CODE_CONFIG, "includes": [
        {"path": "labs/cmdresp/partial/prompt.md", "format": "text"},
        {"path": "docs/labs", "format": "page"}]})
    patterns = config.sparse_patterns()
    assert "/labs/cmdresp/partial/prompt.md" in patterns and "/docs/labs" in patterns
    assert "/materials.xml" in patterns
    assert "!/docs/labs/**/*.png" in patterns


@pytest.mark.parametrize("repo", ["https://gitlab.com/test/hwdesign", "http://github.com/test/hwdesign",
                                  "https://github.com/test", "file:///etc"])
def test_only_a_github_repo_is_accepted(tmp_path, repo) -> None:
    from llmgrader.coursemcp.materials import MaterialsError, read_code

    with pytest.raises(MaterialsError):
        read_code(write_mcp_config(tmp_path, repo=repo))


def test_slide_hyperlinks_are_extracted(tmp_path) -> None:
    pptx = pytest.importorskip("pptx")
    from llmgrader.coursemcp.materials import extract_slides

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "AXI4-Lite"
    run = slide.placeholders[1].text_frame.paragraphs[0].add_run()
    run.text = "see the demo"
    run.hyperlink.address = "https://example.github.io/hwdesign/docs/demos/procif/"
    # A grouped shape: python-pptx refuses a group's click_action outright.
    from pptx.util import Inches
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.click_action.hyperlink.address = "https://github.com/test/hwdesign/tree/main/demos/fsm"
    path = tmp_path / "deck.pptx"
    deck.save(path)
    [extracted] = extract_slides(path)
    assert extracted["links"] == ["https://example.github.io/hwdesign/docs/demos/procif/",
                                  "https://github.com/test/hwdesign/tree/main/demos/fsm"]
    assert extracted["text"] == "see the demo"


# ---------------------------------------------------------------------------
# The real repo
# ---------------------------------------------------------------------------

HWDESIGN_CONFIG = {"repo": "https://github.com/sdrangan/hwdesign", "includes": [
    {"path": "demos"}, {"path": "docs/demos", "format": "page"},
    {"path": "labs"}, {"path": "docs/labs", "format": "page"}]}


@pytest.mark.skipif(not (HWDESIGN / ".git").is_dir(), reason="no hwdesign checkout beside this repo")
def test_the_real_hwdesign_repo() -> None:
    from llmgrader.coursemcp.catalog import pattern_matches
    from llmgrader.coursemcp.code import find_lines

    tracked = set(subprocess.run(["git", "ls-files"], cwd=HWDESIGN, capture_output=True,
                                 text=True, check=True).stdout.splitlines())
    snapshot = build_snapshot(HWDESIGN, CodeConfig.from_dict(HWDESIGN_CONFIG), "HEAD", only=tracked)
    assert not [p for p in snapshot.files if p.endswith((".bit", ".hwh", ".png"))]
    assert "demos/scalar_fun/scalar_fun_rfsoc42/scripts/project_recreate.tcl" not in snapshot.files
    hits = [(s.path, b) for s, b in find_lines(snapshot.files.values(), "s_axilite")
            if s.path == "demos/scalar_fun/scalar_fun_vitis/src/scalar_fun.cpp"]
    assert [(b[0] + 1, b[-1] + 1) for _, b in hits] == [(12, 16)]
    # The labs pair with their docs; the take-home prompt is a text file.
    labs = pattern_matches(snapshot, "labs/*")
    assert {"cmdresp", "cubic", "prng", "rootsolve", "subc"} <= set(labs)
    assert snapshot.files["labs/cmdresp/partial/prompt.md"].format == "text"
    assert {"cmdresp", "cubic"} <= set(pattern_matches(snapshot, "docs/labs/*"))


@pytest.mark.skipif(os.environ.get("LLMGRADER_RUN_NETWORK_TESTS") != "1",
                    reason="clones from GitHub; set LLMGRADER_RUN_NETWORK_TESTS=1")
def test_cloning_the_real_repo_skips_the_bitstreams(tmp_path) -> None:
    config = CodeConfig.from_dict(HWDESIGN_CONFIG)
    sync = CodeSync("hwdesign", config, tmp_path / "code", ttl=10_000)
    assert sync.sync_now()
    assert sync.snapshot.files
    on_disk = [p for p in sync.mirror.rglob("*") if p.is_file() and ".git" not in p.parts]
    assert not [p for p in on_disk if p.suffix in (".bit", ".hwh", ".png")]
    pack_bytes = sum(p.stat().st_size for p in (sync.mirror / ".git").rglob("*") if p.is_file())
    assert pack_bytes < 5_000_000      # the 38 MB of bitstreams were never fetched


@pytest.mark.skipif(os.environ.get("LLMGRADER_RUN_NETWORK_TESTS") != "1",
                    reason="clones from GitHub; set LLMGRADER_RUN_NETWORK_TESTS=1")
def test_the_real_repo_reports_its_commit_and_head(tmp_path) -> None:
    """What Manage Courses shows (plans/demo_code_admin.md), against GitHub:
    a shallow partial clone still has the commit's date and subject, and the
    remote check reads the same head the sync fetched."""
    config = CodeConfig.from_dict(HWDESIGN_CONFIG)
    sync = CodeSync("hwdesign", config, tmp_path / "code", ttl=10_000)
    assert sync.sync_now()
    state = sync.state()
    assert state["subject"] and state["committed_at"][:2] == "20"
    assert sync.remote()["head"] == state["commit"]
    assert sync.check_remote_in_background()
    sync._check_thread.join(60)
    remote = sync.remote()
    assert remote["error"] is None and len(remote["head"]) == 40
    assert sync.sync_now()                       # the ls-remote path: nothing new, or a fetch
    overview = sync.overview()
    assert overview["status"] in ("up_to_date", "update_available")
    assert overview["progress"] is None and not sync.progress_path.exists()
