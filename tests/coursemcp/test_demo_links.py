"""The demo link graph and the build side: ``plans/demo_code_mcp.md``, decision 3.

Pure functions over a directory and slide data -- no git, no app -- plus the
``<code>`` element through the build, slide hyperlinks, a check against the
real hwdesign checkout when one sits beside this repo, and an opt-in clone
from GitHub.
"""

from __future__ import annotations

import os
import subprocess
from pathlib import Path

import pytest

from llmgrader.coursemcp.code import CodeConfig, CodeSync, build_snapshot
from llmgrader.coursemcp.demo_links import (
    OverridesError,
    SlideRef,
    build_index,
    find_references,
    local_report,
    resolve,
    slides_from_materials,
)
from test_course_code import CODE_CONFIG, REPO_BINARIES, REPO_FILES

REPO_ROOT = Path(__file__).resolve().parents[2]
HWDESIGN = REPO_ROOT.parent / "hwdesign"
HWDESIGN_PACKAGE = REPO_ROOT.parent / "hwdesign-soln" / "soln_package"

SLIDES = [
    SlideRef("procif", "Unit 4:  Memory and Processor Interfaces", 50,
             "See demos/scalar_fun/scalar_fun_vitis"),
    SlideRef("fsm", "Unit 2:  Sequential Logic and FSMs", 35, "Run it",
             ("https://example.github.io/hwdesign/docs/demos/stream/",)),
    SlideRef("intro", "Unit 0:  Introduction", 6, "Course site: https://example.github.io/hwdesign/docs"),
]


def make_tree(root: Path, extra: dict | None = None) -> Path:
    tree = root / "tree"
    for path, text in {**REPO_FILES, **(extra or {})}.items():
        (tree / path).parent.mkdir(parents=True, exist_ok=True)
        (tree / path).write_text(text, encoding="utf-8")
    for path, data in REPO_BINARIES.items():
        (tree / path).parent.mkdir(parents=True, exist_ok=True)
        (tree / path).write_bytes(data)
    return tree


def snapshot_of(tmp_path, extra=None):
    return build_snapshot(make_tree(tmp_path, extra), CodeConfig.from_dict(CODE_CONFIG), "abc123")


def index_of(tmp_path, extra=None, slides=SLIDES):
    return build_index(snapshot_of(tmp_path, extra), slides)


# ---------------------------------------------------------------------------
# References
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("text, expected", [
    ("files are in `hwdesign/demos/stream/avgfilt`.", ("dir", "demos/stream/avgfilt")),
    ("see demos/fsm/counter.sv", ("file", "demos/fsm/counter.sv")),
    ("https://github.com/test/hwdesign/blob/main/demos/fsm/counter.sv", ("file", "demos/fsm/counter.sv")),
    ("https://github.com/test/hwdesign/tree/main/demos/stream/poly", ("dir", "demos/stream/poly")),
    ("https://example.github.io/hwdesign/docs/demos/procif/", ("file", "docs/demos/procif/index.md")),
    ("https://example.github.io/hwdesign/docs/demos/procif/vitis_ip.html",
     ("file", "docs/demos/procif/vitis_ip.md")),
    ("see demos/basic_logic/gone.sv", ("broken", "demos/basic_logic/gone.sv")),
    ("demos/scalar_fun/scalar_fun_vitis/hls_component/impl/ip",
     ("build_output", "demos/scalar_fun/scalar_fun_vitis/hls_component/impl/ip")),
    ("demos/scalar_fun/overlay/design.bit", ("excluded", "demos/scalar_fun/overlay/design.bit")),
    ("https://example.github.io/hwdesign/docs", ("outside", "docs")),
])
def test_each_reference_form_resolves(tmp_path, text, expected) -> None:
    snapshot = snapshot_of(tmp_path)
    [(_, candidates)] = find_references(text, snapshot)
    assert resolve(candidates, snapshot) == expected


def test_a_path_inside_a_longer_path_is_not_a_reference(tmp_path) -> None:
    snapshot = snapshot_of(tmp_path)
    assert find_references("Wrote C:/work/demos/stream/avgfilt/out.csv", snapshot) == []


def test_relative_links_between_docs_pages(tmp_path) -> None:
    snapshot = snapshot_of(tmp_path)
    [(_, candidates)] = find_references("the [example](../procif/) first", snapshot,
                                        page="docs/demos/stream/index.md")
    assert resolve(candidates, snapshot) == ("file", "docs/demos/procif/index.md")


# ---------------------------------------------------------------------------
# Demos
# ---------------------------------------------------------------------------


def test_demos_and_their_links_are_derived(tmp_path) -> None:
    index = index_of(tmp_path)
    demos = index.demos
    assert sorted(demos) == ["fifoif", "fsm", "procif", "stream"]
    assert demos["procif"].documented and demos["procif"].code == ["demos/scalar_fun/scalar_fun_vitis"]
    assert demos["stream"].code == ["demos/stream/avgfilt"]
    assert demos["fifoif"].code == ["demos/stream/poly"]
    assert not demos["fsm"].documented and demos["fsm"].code == ["demos/fsm"]
    assert demos["procif"].slides == [{"deck": "procif", "slide": 50,
                                       "via": "slide procif:50 cites demos/scalar_fun/scalar_fun_vitis"}]
    assert [u["unit"] for u in demos["stream"].units] == ["Unit 2:  Sequential Logic and FSMs"]
    assert demos["fsm"].units[0]["via"] == "deck fsm has the demo's name"
    # A site root (the intro slide) names no demo.
    assert not any(s["deck"] == "intro" for d in demos.values() for s in d.slides)
    assert index.demos_for_path("demos/scalar_fun/scalar_fun_vitis/src/scalar_fun.cpp") == ["procif"]
    assert [d.id for d in index.demos_for_slide("procif", 50)] == ["procif"]
    assert index.doc_cites["docs/demos/procif/index.md"] == ["demos/scalar_fun/scalar_fun_vitis"]


def test_the_same_name_directory_belongs_to_its_docs_topic(tmp_path) -> None:
    index = index_of(tmp_path, {"docs/demos/fsm/index.md": "---\ntitle: FSM demo\n---\n\nCounters.\n"})
    fsm = index.demos["fsm"]
    assert fsm.documented and fsm.code == ["demos/fsm"] and fsm.title == "FSM demo"
    assert "fsm-code" not in index.demos


def test_broken_and_build_output_references_are_reported(tmp_path) -> None:
    report = index_of(tmp_path).report
    assert report.broken == [("docs/demos/procif/vitis_ip.md:17", "demos/basic_logic/old_fun.sv")]
    lines = "\n".join(report.lines(index_of(tmp_path)))
    assert "Broken reference" in lines and "demos/basic_logic/old_fun.sv" in lines


def test_a_code_only_demo_takes_its_readme(tmp_path) -> None:
    index = index_of(tmp_path, {"demos/fsm/README.md": "# The counter demo\n\nA counting FSM.\n"})
    assert index.demos["fsm"].title == "The counter demo"
    assert index.demos["fsm"].description == "A counting FSM."


# ---------------------------------------------------------------------------
# Overrides
# ---------------------------------------------------------------------------


OVERRIDES = """<demos>
  <demo id="fsm" title="Counter FSM">
    <description>A counter, written as an FSM.</description>
    <unit>Unit 9:  Extra</unit>
    <slides deck="procif" first="3" last="4"/>
  </demo>
  <demo id="procif"><not_slides deck="procif"/></demo>
  <hide id="fifoif"/>
  <demo id="nosuch"><unit>Unit 1</unit></demo>
  <demo id="stream"><slides deck="nodeck"/></demo>
</demos>
"""


def test_overrides_fill_in_correct_and_hide(tmp_path) -> None:
    index = index_of(tmp_path, {"demos/demos.xml": OVERRIDES})
    fsm = index.demos["fsm"]
    assert fsm.title == "Counter FSM" and fsm.description == "A counter, written as an FSM."
    assert {u["unit"] for u in fsm.units} >= {"Unit 9:  Extra", "Unit 4:  Memory and Processor Interfaces"}
    assert [(s["deck"], s["slide"]) for s in fsm.slides] == [("procif", 3), ("procif", 4)]
    assert index.demos["procif"].slides == [] and index.demos["procif"].units == []
    assert "fifoif" not in [d.id for d in index.visible()]
    problems = " ".join(index.report.unknown_overrides)
    assert "'nosuch'" in problems and "'nodeck'" in problems


@pytest.mark.parametrize("text", ["<demos><demo id='fsm'>", "<demos><demo id='fsm'><bogus/></demo></demos>",
                                  "<demos><demo/></demos>"])
def test_a_broken_overrides_file_is_an_error(tmp_path, text) -> None:
    with pytest.raises(OverridesError):
        index_of(tmp_path, {"demos/demos.xml": text})


def test_symlinks_are_never_served(tmp_path) -> None:
    tree = make_tree(tmp_path)
    secret = tmp_path / "secret.txt"
    secret.write_text("SECRET", encoding="utf-8")
    try:
        os.symlink(secret, tree / "demos" / "fsm" / "leak.txt")
    except (OSError, NotImplementedError):
        pytest.skip("this system cannot create symlinks")
    snapshot = build_snapshot(tree, CodeConfig.from_dict(CODE_CONFIG), "abc")
    assert "demos/fsm/leak.txt" not in snapshot.files
    assert snapshot.search("SECRET")["hits"] == []


# ---------------------------------------------------------------------------
# The build: <code>, slide hyperlinks, the report
# ---------------------------------------------------------------------------


MCP_CONFIG = """<llmgrader_mcp>
  <roots><root id="pub" path="{tree}"/></roots>
  <code repo="{repo}" root="pub">
    <include path="demos"/>
    <include path="docs/demos" kind="doc"/>
    <exclude>*.wcfg</exclude>
  </code>
</llmgrader_mcp>
"""


def write_mcp_config(tmp_path, repo="https://github.com/test/hwdesign") -> Path:
    tree = make_tree(tmp_path)
    path = tmp_path / "llmgrader_mcp_config.xml"
    path.write_text(MCP_CONFIG.format(tree=tree.as_posix(), repo=repo), encoding="utf-8")
    return path


def test_code_is_read_from_the_config_and_carried_in_the_manifest(tmp_path) -> None:
    from llmgrader.coursemcp.materials import build_materials, load_materials, read_code

    spec = read_code(write_mcp_config(tmp_path))
    assert spec.config["repo"] == "https://github.com/test/hwdesign"
    assert spec.config["includes"] == [{"path": "demos", "kind": "code"},
                                       {"path": "docs/demos", "kind": "doc"}]
    assert spec.config["excludes"] == ["*.wcfg"]
    assert spec.local_root == (tmp_path / "tree").resolve()
    package = tmp_path / "pkg"
    build_materials([], package, code=spec.config, log=lambda *_: None)
    assert load_materials(str(package)).code == spec.config


@pytest.mark.parametrize("repo", ["https://gitlab.com/test/hwdesign", "http://github.com/test/hwdesign",
                                  "https://github.com/test", "file:///etc"])
def test_only_a_github_repo_is_accepted(tmp_path, repo) -> None:
    from llmgrader.coursemcp.materials import MaterialsError, read_code

    with pytest.raises(MaterialsError):
        read_code(write_mcp_config(tmp_path, repo=repo))


def test_the_link_report_reads_a_local_checkout(tmp_path) -> None:
    from llmgrader.coursemcp.materials import read_code

    spec = read_code(write_mcp_config(tmp_path))
    lines = "\n".join(local_report(spec.config, spec.local_root, None))
    assert "no slides" in lines and "procif" in lines
    assert "Broken reference: docs/demos/procif/vitis_ip.md:17" in lines


def test_slide_hyperlinks_are_extracted(tmp_path) -> None:
    pptx = pytest.importorskip("pptx")
    from llmgrader.coursemcp.materials import extract_slides

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "AXI4-Lite"
    run = slide.placeholders[1].text_frame.paragraphs[0].add_run()
    run.text = "see the demo"
    run.hyperlink.address = "https://example.github.io/hwdesign/docs/demos/procif/"
    # A grouped shape: python-pptx refuses a group's click_action outright.
    from pptx.util import Inches
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.click_action.hyperlink.address = "https://github.com/test/hwdesign/tree/main/demos/fsm"
    path = tmp_path / "deck.pptx"
    deck.save(path)
    [extracted] = extract_slides(path)
    assert extracted["links"] == ["https://example.github.io/hwdesign/docs/demos/procif/",
                                  "https://github.com/test/hwdesign/tree/main/demos/fsm"]
    assert extracted["text"] == "see the demo"


# ---------------------------------------------------------------------------
# The real repo
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not (HWDESIGN / ".git").is_dir(), reason="no hwdesign checkout beside this repo")
def test_the_real_hwdesign_demos() -> None:
    from llmgrader.coursemcp.materials import load_materials

    config = CodeConfig.from_dict({"repo": "https://github.com/sdrangan/hwdesign",
                                   "includes": [{"path": "demos"}, {"path": "docs/demos", "kind": "doc"}]})
    tracked = set(subprocess.run(["git", "ls-files"], cwd=HWDESIGN, capture_output=True,
                                 text=True, check=True).stdout.splitlines())
    snapshot = build_snapshot(HWDESIGN, config, "HEAD", only=tracked)
    assert not [p for p in snapshot.files if p.endswith((".bit", ".hwh", ".png"))]
    assert "demos/scalar_fun/scalar_fun_rfsoc42/scripts/project_recreate.tcl" not in snapshot.files
    materials = load_materials(str(HWDESIGN_PACKAGE)) if HWDESIGN_PACKAGE.is_dir() else None
    index = build_index(snapshot, slides_from_materials(materials))
    hits = [h for h in snapshot.search("s_axilite", demos_for=index.demos_for_path)["hits"]
            if h["path"] == "demos/scalar_fun/scalar_fun_vitis/src/scalar_fun.cpp"]
    assert [(h["start_line"], h["end_line"], h["demos"]) for h in hits] == [(12, 16, ["procif"])]
    if materials is not None:
        assert [d.id for d in index.demos_for_slide("procif", 50)] == ["procif"]
        assert [d.id for d in index.demos_for_slide("fifo", 28)] == ["stream"]
        assert [d.id for d in index.demos_for_slide("fifo", 55)] == ["fifoif"]
        assert [d.id for d in index.demos_for_slide("fsm", 35)] == ["fsm"]


@pytest.mark.skipif(os.environ.get("LLMGRADER_RUN_NETWORK_TESTS") != "1",
                    reason="clones from GitHub; set LLMGRADER_RUN_NETWORK_TESTS=1")
def test_cloning_the_real_repo_skips_the_bitstreams(tmp_path) -> None:
    config = CodeConfig.from_dict({"repo": "https://github.com/sdrangan/hwdesign",
                                   "includes": [{"path": "demos"}, {"path": "docs/demos", "kind": "doc"}]})
    sync = CodeSync("hwdesign", config, tmp_path / "code", ttl=10_000)
    assert sync.sync_now()
    assert sync.snapshot.files
    on_disk = [p for p in sync.mirror.rglob("*") if p.is_file() and ".git" not in p.parts]
    assert not [p for p in on_disk if p.suffix in (".bit", ".hwh", ".png")]
    pack_bytes = sum(p.stat().st_size for p in (sync.mirror / ".git").rglob("*") if p.is_file())
    assert pack_bytes < 5_000_000      # the 38 MB of bitstreams were never fetched
