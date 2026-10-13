"""Course MCP material built into the package, and the package's half of it served.

``plans/unified_materials.md``: a course declares its own kinds of material
(slides, demo, lab, handout...), units that hold only material, and items,
in ``<materials>`` of ``llmgrader_mcp_config.xml``.  An item comes from one
of two sources.  This module owns the configuration and the **package**
source; the **mirror** source is ``code.py``'s, and ``catalog.py`` joins the
two into the one index the tools read.

It also carries the units published to the MCP alone -- past exams, which
the portal must not list (``plans/exam_units.md``) -- and a description of
each kind of unit.

**Build** (instructor's machine, ``llmgrader_mcp_build`` or ``create_soln_pkg``):
a package item's files are read by their extension -- a ``.pptx`` gives every
slide's title, text and speaker notes and its PDF export each slide's image;
a ``.pdf`` gives each page's text and image; anything else is copied as text
-- and written under ``<package>/mcp_materials/``, so the course package --
one upload -- carries it.  Package paths are explicit, never patterns.
python-pptx and PyMuPDF are needed here only, and are imported lazily.

**Serve** (the portal): plain JSON, text and JPEGs read from the loaded
package.  No new dependencies on the server, and no embeddings: the whole
course is ~500 slides, and the AI calling the tool rephrases and searches
again on its own -- the synonym work embeddings would otherwise do
(``plans/course_mcp.md``, decision 2).

Most of a lecture slide's content is in its figures, not its text, which is
why a slide is served as its image as well.
"""

from __future__ import annotations

import fnmatch
import json
import math
import os
import posixpath
import re
import threading
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

MATERIALS_DIR = "mcp_materials"
MANIFEST = "manifest.json"
FORMAT_VERSION = 2
IMAGE_WIDTH_PX = 1280
JPEG_QUALITY = 80
SCHEMA_PATH = Path(__file__).resolve().parents[1] / "schemas" / "llmgrader_mcp_config.xsd"

# How a file is read: plans/unified_materials.md, "Format".  It follows from
# the file, never from the kind.
SLIDES, PDF, PAGE, TEXT = "slides", "pdf", "page", "text"
PAGED_FORMATS = (SLIDES, PDF)
POLICIES = ("open", "tutor", "private")


class MaterialsError(Exception):
    """A configuration or build problem, with a message for the instructor."""


# ---------------------------------------------------------------------------
# Units: one resolver for the build and the server
# ---------------------------------------------------------------------------


def _norm(text: str) -> str:
    return " ".join((text or "").lower().split())


class UnitResolver:
    """Turns a unit reference into the unit's name, or None.

    A reference is a unit ``<name>`` -- from llmgrader_config.xml, an
    MCP-only unit, or a material ``<unit>`` -- or a material unit's id.  Names
    match whatever their case and spacing: unit names carry double spaces
    ("Unit 2:  Sequential Logic"), which nobody types twice the same way.
    """

    def __init__(self, names: list[str], material_units: list[dict] | None = None):
        self.names = list(names)
        self.material_units = list(material_units or [])
        self._by_id = {u["id"]: u["name"] for u in self.material_units}
        self._by_name = {}
        for name in [*self.names, *(u["name"] for u in self.material_units)]:
            self._by_name.setdefault(_norm(name), name)

    def resolve(self, ref: str | None) -> str | None:
        ref = (ref or "").strip()
        if not ref:
            return None
        if ref in self._by_id:
            return self._by_id[ref]
        return self._by_name.get(_norm(ref))

    def all_names(self) -> list[str]:
        return list(self._by_name.values())


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------


@dataclass
class DeckSpec:
    """One paged document in the package: a slide deck (``pptx``, with its
    PDF export for images) or a PDF read page by page (``pptx`` None)."""
    id: str                    # the item's id, or <item>-<n> for an item with several
    unit: str                  # the item's first unit: what the describe prompt names
    title: str
    pptx: Path | None
    pdf: Path | None
    item: str = ""
    path: str = ""             # as the config writes it: what the assistant is shown

    @property
    def format(self) -> str:
        return SLIDES if self.pptx is not None else PDF


@dataclass
class PartSpec:
    """One file or folder of an item, as the config names it."""
    path: str                  # relative to its root (package) or the repo (mirror)
    source_path: Path | None = None    # package: the file or folder on disk
    pdf: Path | None = None            # package: a .pptx's PDF export
    docs: bool = False                 # mirror: a docs page or folder


@dataclass
class ItemSpec:
    id: str
    kind: str
    source: str                # "package" or "mirror"
    title: str = ""
    description: str = ""
    units: list[str] = field(default_factory=list)      # resolved names
    parts: list[PartSpec] = field(default_factory=list)
    slides: list[dict] = field(default_factory=list)    # <slides deck first last>
    not_slides: list[str] = field(default_factory=list)


@dataclass
class MaterialsSpec:
    """``<materials>``, read and checked."""
    kinds: list[dict] = field(default_factory=list)            # {id, description, policy}
    material_units: list[dict] = field(default_factory=list)   # {id, name, description}
    items: list[ItemSpec] = field(default_factory=list)        # explicit items, both sources
    patterns: list[dict] = field(default_factory=list)         # {kind, path, docs, units}
    refinements: list[dict] = field(default_factory=list)      # for items the build cannot see
    decks: list[DeckSpec] = field(default_factory=list)

    def kind(self, kind_id: str) -> dict | None:
        return next((k for k in self.kinds if k["id"] == kind_id), None)

    def policy(self, kind_id: str) -> str:
        found = self.kind(kind_id)
        return found["policy"] if found else "open"


@dataclass
class UnitSpec:
    """A unit published to the course MCP only (plans/exam_units.md, decision 2)."""
    name: str
    section: str
    path: Path            # the unit XML
    images: Path | None   # the images/ folder beside it, if any

    @property
    def stem(self) -> str:
        return self.path.stem


def _load_config(config_path: str | Path, root_overrides: dict[str, str] | None):
    """``(config_path, root element, {root id: path})``, validated."""
    import xml.etree.ElementTree as ET

    import xmlschema

    config_path = Path(config_path).resolve()
    if not config_path.is_file():
        raise MaterialsError(
            f"No {config_path.name} in {config_path.parent}.\n"
            f"Run this from the folder that holds your course's llmgrader_config.xml "
            f"and {config_path.name}, or pass --config <path>.")
    try:
        ET.parse(config_path)
    except ET.ParseError as exc:
        raise MaterialsError(f"{config_path.name} is not well-formed XML: {exc}") from exc
    schema = xmlschema.XMLSchema(str(SCHEMA_PATH))
    errors = [str(e.reason or e) for e in schema.iter_errors(str(config_path))]
    if errors:
        raise MaterialsError(f"{config_path.name} is not valid: " + "; ".join(errors))

    root_elem = ET.parse(config_path).getroot()
    base = config_path.parent
    roots = {"": base}
    for elem in root_elem.findall("roots/root"):
        roots[elem.get("id")] = (base / elem.get("path")).resolve()
    for root_id, path in (root_overrides or {}).items():
        if root_id not in roots:
            raise MaterialsError(f"--root {root_id}=...: no <root id={root_id!r}> in {config_path.name}")
        roots[root_id] = Path(path).resolve()
    return config_path, root_elem, roots


def _unit_refs(elem) -> list[str]:
    """Every unit an ``<item>`` or ``<items>`` names: unit=, also=, <unit>."""
    refs = [elem.get("unit"), elem.get("also")]
    refs += [u.text for u in elem.findall("unit")]
    return [r.strip() for r in refs if r and r.strip()]


def refinement_of(elem) -> dict:
    """An ``<item>`` as a refinement: the fields it sets, units unresolved.

    The same shape for the config's ``<item id>`` and the repo's
    ``materials.xml``, which is why the server can apply both alike.
    """
    description = elem.findtext("description")
    return {
        "id": elem.get("id"),
        "kind": elem.get("kind"),
        "title": (elem.get("title") or "").strip() or None,
        "description": " ".join(description.split()) if description and description.strip() else None,
        "units": _unit_refs(elem),
        "slides": [{"deck": s.get("deck"),
                    "first": int(s.get("first")) if s.get("first") else None,
                    "last": int(s.get("last")) if s.get("last") else None}
                   for s in elem.findall("slides")],
        "not_slides": [s.get("deck") for s in elem.findall("not_slides")],
    }


def _is_glob(path: str) -> bool:
    return any(ch in path for ch in "*?[")


def portal_unit_names(config_path: str | Path) -> list[str]:
    """The unit names llmgrader_config.xml beside *config_path* lists."""
    return list(_portal_units(Path(config_path).resolve()).values())


def read_materials(config_path: str | Path, root_overrides: dict[str, str] | None = None,
                   *, unit_names: list[str] | None = None) -> MaterialsSpec:
    """``<materials>``, with every package path resolved and checked.

    *unit_names* are the units an item may name besides the material units:
    by default the portal's (llmgrader_config.xml) and the MCP-only ones.
    Refinements of an item made here are applied here; those of an item only
    a mirror pattern can make are kept for the server, and checked by
    :func:`check_catalog` when there is a local checkout to check them on.
    """
    config_path, root_elem, roots = _load_config(config_path, root_overrides)
    spec = MaterialsSpec()
    elem = root_elem.find("materials")
    if elem is None:
        return spec
    problems: list[str] = []
    for kind in elem.findall("kind"):
        spec.kinds.append({"id": kind.get("id"),
                           "description": " ".join(kind.get("description").split()),
                           "policy": kind.get("policy") or "open"})
    for unit in elem.findall("unit"):
        spec.material_units.append({"id": unit.get("id"), "name": unit.get("name").strip(),
                                    "description": " ".join((unit.get("description") or "").split())})
    if unit_names is None:
        unit_names = portal_unit_names(config_path)
        try:
            unit_names += [u.name for u in read_units(config_path, root_overrides)]
        except MaterialsError:
            pass        # read_units' own call reports it
    resolver = UnitResolver(unit_names, spec.material_units)
    taken = {u["id"] for u in spec.material_units}
    clash = taken & {_norm(n) for n in unit_names}
    if clash:
        problems.append(f"material unit id(s) {', '.join(sorted(clash))} are also unit names")

    def units_of(refs: list[str], where: str) -> list[str]:
        names = []
        for ref in refs:
            name = resolver.resolve(ref)
            if name is None:
                problems.append(f"{where}: unknown unit {ref!r}; units are: "
                                + "; ".join(repr(n) for n in resolver.all_names()))
            elif name not in names:
                names.append(name)
        return names

    kinds = {k["id"] for k in spec.kinds}
    has_code = root_elem.find("code") is not None
    refinements: list = []
    for child in elem:
        if child.tag == "items":
            where = f"<items path={child.get('path')!r}>"
            if child.get("kind") not in kinds:
                problems.append(f"{where}: unknown kind {child.get('kind')!r}")
            if not has_code:
                problems.append(f"{where}: a mirror pattern needs <code repo=...>")
            path, docs = child.get("path"), child.get("docs")
            for pattern in filter(None, [path, docs]):
                head, _, last = pattern.rpartition("/")
                if _is_glob(head):
                    problems.append(f"{where}: only the last part of {pattern!r} may be a pattern")
            spec.patterns.append({"kind": child.get("kind"), "path": path, "docs": docs,
                                  "units": units_of(_unit_refs(child), where)})
        elif child.tag == "item":
            if child.get("path") or child.get("docs") or child.findall("part"):
                item = _read_item(child, roots, has_code, problems)
                if item is not None:
                    if item.kind not in kinds:
                        problems.append(f"item {item.id!r}: unknown kind {item.kind!r}")
                    item.units = units_of(_unit_refs(child), f"item {item.id!r}")
                    spec.items.append(item)
            else:
                refinements.append(child)

    seen: set[str] = set()
    for item in spec.items:
        if item.id in seen:
            problems.append(f"two items have the id {item.id!r}")
        seen.add(item.id)
    by_id = {item.id: item for item in spec.items}
    for child in refinements:
        refinement = refinement_of(child)
        where = f"<item id={refinement['id']!r}>"
        if refinement["kind"] and refinement["kind"] not in kinds:
            problems.append(f"{where}: unknown kind {refinement['kind']!r}")
        refinement["units"] = units_of(refinement["units"], where)
        item = by_id.get(refinement["id"])
        if item is None:
            if not spec.patterns:
                problems.append(f"{where} matches no item, and there is no <items> pattern "
                                "it could refine")
            spec.refinements.append(refinement)
            continue
        item.kind = refinement["kind"] or item.kind
        item.title = refinement["title"] or item.title
        item.description = refinement["description"] or item.description
        item.units += [u for u in refinement["units"] if u not in item.units]
        item.slides += refinement["slides"]
        item.not_slides += refinement["not_slides"]

    missing = [item.id for item in spec.items if item.source == "package" and not item.units]
    if missing:
        problems.append("every item needs a unit; these have none: " + ", ".join(missing))
    if problems:
        raise MaterialsError("; ".join(problems))
    spec.decks = _deck_specs(spec)
    return spec


def _read_item(elem, roots: dict, has_code: bool, problems: list[str]) -> ItemSpec | None:
    item_id = elem.get("id")
    source = elem.get("source") or "package"
    where = f"item {item_id!r}"
    if not elem.get("kind"):
        problems.append(f"{where}: an item with a path needs kind=")
        return None
    description = elem.findtext("description")
    item = ItemSpec(id=item_id, kind=elem.get("kind"), source=source,
                    title=(elem.get("title") or "").strip(),
                    description=" ".join(description.split()) if description else "")
    item.slides = refinement_of(elem)["slides"]
    item.not_slides = refinement_of(elem)["not_slides"]
    raw = []
    if elem.get("path") or elem.get("pdf"):
        raw.append({"path": elem.get("path"), "pdf": elem.get("pdf"), "docs": None})
    if elem.get("docs"):
        raw.append({"path": None, "pdf": None, "docs": elem.get("docs")})
    raw += [{"path": p.get("path"), "pdf": p.get("pdf"), "docs": p.get("docs")}
            for p in elem.findall("part")]

    if source == "mirror":
        if elem.get("root"):
            problems.append(f"{where}: root= is for package items; a mirror item's paths "
                            "are in the <code> repo")
        if not has_code:
            problems.append(f"{where}: source=\"mirror\" needs <code repo=...>")
        for part in raw:
            if part["pdf"]:
                problems.append(f"{where}: pdf= is for a package .pptx")
            for value, docs in ((part["path"], False), (part["docs"], True)):
                if not value:
                    continue
                value = value.strip().strip("/")
                if _is_glob(value) or ".." in value.split("/") or re.match(r"^[A-Za-z]:", value):
                    problems.append(f"{where}: {value!r} is not a plain repo path "
                                    "(patterns belong in <items>)")
                    continue
                item.parts.append(PartSpec(value, docs=docs))
        return item

    root_id = elem.get("root") or ""
    if root_id not in roots:
        problems.append(f"{where}: unknown root {root_id!r}")
        return None
    base = roots[root_id]
    for part in raw:
        if part["docs"]:
            problems.append(f"{where}: docs= is for mirror items; a package item lists "
                            "its pages with path=")
            continue
        relative = (part["path"] or "").strip().replace("\\", "/").strip("/")
        if not relative:
            problems.append(f"{where}: pdf= needs the .pptx it was exported from in path=")
            continue
        if _is_glob(relative):
            problems.append(f"{where}: {relative!r} is a pattern; package paths are explicit "
                            "(patterns are for the mirror, in <items>)")
            continue
        full = (base / relative).resolve()
        if not full.exists():
            problems.append(f"{where}: {full} not found")
            continue
        pdf = None
        if part["pdf"]:
            if full.suffix.lower() != ".pptx":
                problems.append(f"{where}: pdf= goes with a .pptx path")
                continue
            pdf = (base / part["pdf"]).resolve()
            if not pdf.is_file():
                problems.append(f"{where}: pdf {pdf} not found")
                continue
        item.parts.append(PartSpec(relative, source_path=full, pdf=pdf))
    return item


def _deck_specs(spec: MaterialsSpec) -> list[DeckSpec]:
    """The paged documents of the package items: what the build renders and
    describe describes.  A private item has none: it is never built."""
    decks = []
    for item in spec.items:
        if item.source != "package" or spec.policy(item.kind) == "private":
            continue
        paged = []
        for part in item.parts:
            if part.source_path.is_dir():
                paged += [(f"{part.path}/{f.relative_to(part.source_path).as_posix()}", None, f)
                          for f in sorted(part.source_path.rglob("*"))
                          if f.is_file() and f.suffix.lower() == ".pdf"
                          and not _skipped(f.relative_to(part.source_path).as_posix())]
            elif part.source_path.suffix.lower() == ".pptx":
                paged.append((part.path, part.source_path, part.pdf))
            elif part.source_path.suffix.lower() == ".pdf":
                paged.append((part.path, None, part.source_path))
        for number, (path, pptx, pdf) in enumerate(paged, start=1):
            deck_id = item.id if len(paged) == 1 else f"{item.id}-{number}"
            # A folder part's PDFs arrive as (path, None, file).
            decks.append(DeckSpec(deck_id, item.units[0] if item.units else "",
                                  item.title if len(paged) == 1 else "",
                                  pptx, pdf, item=item.id, path=path))
    return decks


def _skipped(relative: str) -> bool:
    from llmgrader.coursemcp.code import is_build_output
    return is_build_output(relative)


def _portal_units(config_path: Path) -> dict[Path, str]:
    """The units llmgrader_config.xml beside *config_path* lists: source -> name."""
    import xml.etree.ElementTree as ET

    portal_config = config_path.parent / "llmgrader_config.xml"
    try:
        root = ET.parse(portal_config).getroot()
    except (OSError, ET.ParseError):
        return {}
    units = {}
    for elem in root.findall("units/unit"):
        source, name = elem.findtext("source"), elem.findtext("name")
        if source and name:
            units[(portal_config.parent / source.strip()).resolve()] = name.strip()
    return units


def read_units(config_path: str | Path, root_overrides: dict[str, str] | None = None) -> list[UnitSpec]:
    """The MCP-only units *config_path* lists, each checked against unit.xsd.

    A unit that llmgrader_config.xml also lists is refused: this list is for
    units the portal does not have, and one in both would be served twice.
    So is a name the portal already uses, since tools look units up by name.
    """
    import xml.etree.ElementTree as ET

    from llmgrader.services.unit_parser import UnitParser

    config_path, root_elem, roots = _load_config(config_path, root_overrides)
    portal = _portal_units(config_path)
    specs, problems = [], []
    seen_names: set[str] = set()
    seen_stems: dict[str, Path] = {}
    for elem in root_elem.findall("units/unit"):
        relative = (elem.text or "").strip()
        root_id = elem.get("root") or ""
        if root_id not in roots:
            problems.append(f"unit {relative!r}: unknown root {root_id!r}")
            continue
        path = (roots[root_id] / relative).resolve()
        if not path.is_file():
            problems.append(f"unit {relative!r}: {path} not found")
            continue
        if path in portal:
            problems.append(f"unit {relative!r} is also in llmgrader_config.xml (as "
                            f"{portal[path]!r}); list it in one config or the other")
            continue
        errors = UnitParser.validate_unit_file(str(path))
        if errors:
            problems.append(f"unit {relative!r} is not valid: " + "; ".join(errors))
            continue
        unit_root = ET.parse(path).getroot()
        name = (elem.get("name") or unit_root.get("title") or path.stem).strip()
        if name in portal.values():
            problems.append(f"unit {relative!r}: the portal already has a unit named {name!r}; "
                            "give this one a different name=")
            continue
        if name in seen_names:
            problems.append(f"unit {relative!r}: another unit is already named {name!r}")
            continue
        # The stem names the unit's file and its images folder in the package.
        if path.stem in seen_stems:
            problems.append(f"unit {relative!r}: same file name as {seen_stems[path.stem]}; "
                            "rename one")
            continue
        seen_names.add(name)
        seen_stems[path.stem] = path
        images = path.parent / "images"
        specs.append(UnitSpec(name, (elem.get("section") or "").strip(), path,
                              images if images.is_dir() else None))
    if problems:
        raise MaterialsError("; ".join(problems))
    return specs


def read_unit_types(config_path: str | Path) -> list[dict]:
    """The unit type descriptions *config_path* gives: id, title, description."""
    _, root_elem, _ = _load_config(config_path, None)
    return [
        {"id": elem.get("id"), "title": (elem.get("title") or "").strip(),
         "description": " ".join((elem.text or "").split())}
        for elem in root_elem.findall("unit_types/unit_type")
    ]


@dataclass
class CodeSpec:
    """``<code>`` as the manifest carries it, and the local checkout, if a
    ``root=`` names one, for the build's check of the mirrored items."""
    config: dict
    local_root: Path | None


def mirror_includes(spec: MaterialsSpec) -> list[dict]:
    """What the portal must fetch for *spec*'s mirror items and patterns:
    each path as an include, ``page`` for docs and ``text`` otherwise, with
    any path inside another of its format dropped."""
    wanted: list[tuple[str, str]] = []
    for item in spec.items:
        if item.source == "mirror":
            wanted += [(part.path, "page" if part.docs else "text") for part in item.parts]
    for pattern in spec.patterns:
        wanted.append((pattern["path"].rpartition("/")[0] or pattern["path"], "text"))
        if pattern["docs"]:
            wanted.append((pattern["docs"].rpartition("/")[0] or pattern["docs"], "page"))
    includes = []
    for path, fmt in wanted:
        if any(fmt == f and (path == p or path.startswith(p + "/")) for p, f in wanted
               if (p, f) != (path, fmt)):
            continue
        if {"path": path, "format": fmt} not in includes:
            includes.append({"path": path, "format": fmt})
    return includes


def read_code(config_path: str | Path, root_overrides: dict[str, str] | None = None,
              *, spec: MaterialsSpec | None = None) -> CodeSpec | None:
    """The mirror *config_path* names (``<code>``), or None.

    Only a pointer: the files themselves are synced by the portal, never
    packaged (``plans/demo_code_mcp.md``, decision 2).  Which paths it
    fetches follows from the mirror items and patterns in *spec*.
    """
    _, root_elem, roots = _load_config(config_path, root_overrides)
    elem = root_elem.find("code")
    if elem is None:
        return None
    root_id = elem.get("root")
    if root_id and root_id not in roots:
        raise MaterialsError(f"<code root={root_id!r}>: no <root id={root_id!r}>")
    if spec is None:
        spec = read_materials(config_path, root_overrides)
    config = {
        "repo": elem.get("repo").rstrip("/").removesuffix(".git"),
        "branch": elem.get("branch") or "main",
        "includes": mirror_includes(spec),
        "excludes": [(e.text or "").strip() for e in elem.findall("exclude") if (e.text or "").strip()],
        "site": (elem.get("site") or "").rstrip("/") or None,
        "overrides": elem.get("overrides") or None,
    }
    return CodeSpec(config, roots[root_id] if root_id else None)


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


# PP_PLACEHOLDER: DATE 16, FOOTER 15, SLIDE_NUMBER 13.  Their text is the
# deck's chrome ("35", "Fall 2026"), not the slide's content: left in, a
# figure-only slide looks as if it had text.
_CHROME_PLACEHOLDERS = {13, 15, 16}


def _is_chrome(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return int(shape.placeholder_format.type) in _CHROME_PLACEHOLDERS
    except (AttributeError, ValueError, TypeError):
        return False


def _shape_texts(shape) -> list[str]:
    """Every piece of text in *shape*: text frames, tables, grouped shapes."""
    texts = []
    if _is_chrome(shape):
        return texts
    if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
        for child in shape.shapes:
            texts.extend(_shape_texts(child))
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            texts.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                texts.append(" | ".join(cells))
    return texts


def _clean(text: str) -> str:
    # PowerPoint writes a soft line break inside a paragraph as a vertical tab.
    return text.replace("\x0b", " ").strip()


def _shape_links(shape) -> list[str]:
    """Hyperlink targets in *shape*: its click action and every text run's link.

    A slide often says "see the demo" with the address only in the link, so
    the visible text alone would miss it (``plans/demo_code_mcp.md``).
    """
    links = []
    if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
        # A group has no click action or text of its own (python-pptx raises
        # TypeError for its click_action): only its children do.
        for child in shape.shapes:
            links.extend(_shape_links(child))
        return links
    try:
        address = shape.click_action.hyperlink.address
        if address:
            links.append(address)
    except (AttributeError, KeyError, ValueError, TypeError):
        pass
    frames = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        frames.append(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
    for frame in frames:
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                try:
                    address = run.hyperlink.address
                except (AttributeError, KeyError, ValueError):
                    address = None
                if address:
                    links.append(address)
    return links


def extract_slides(pptx_path: Path) -> list[dict]:
    """Title, body text, speaker notes and hyperlinks of every slide, in order."""
    from pptx import Presentation

    slides = []
    for number, slide in enumerate(Presentation(str(pptx_path)).slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        body = []
        links: list[str] = []
        for shape in slide.shapes:
            for address in _shape_links(shape):
                if address not in links:
                    links.append(address)
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            body.extend(_shape_texts(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if not title and body:
            title = body[0].splitlines()[0][:80]
        slides.append({"n": number, "title": " ".join(_clean(title).split()),
                       "text": _clean("\n".join(body)), "notes": _clean(notes),
                       "links": links})
    return slides


def render_pages(pdf_path: Path, out_dir: Path) -> list[str]:
    """Each PDF page as a JPEG in *out_dir*; returns the file names, in order."""
    import pymupdf

    names = []
    with pymupdf.open(str(pdf_path)) as doc:
        for index, page in enumerate(doc, start=1):
            zoom = IMAGE_WIDTH_PX / page.rect.width
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            name = f"slide-{index:03d}.jpg"
            (out_dir / name).write_bytes(pix.tobytes("jpeg", jpg_quality=JPEG_QUALITY))
            names.append(name)
    return names


def pdf_page_count(pdf_path: Path) -> int:
    import pymupdf

    with pymupdf.open(str(pdf_path)) as doc:
        return doc.page_count


def image_key(data: bytes) -> str:
    """A slide image's identity for the description cache: its content hash.

    Content, not deck and number: inserting a slide renumbers every one after
    it, and re-exporting an unchanged deck must not pay to describe it again.
    """
    import hashlib
    return hashlib.sha256(data).hexdigest()


def _page_title(text: str) -> str:
    for line in (text or "").splitlines():
        if line.strip():
            return " ".join(line.split())[:80]
    return ""


def pdf_pages(pdf_path: Path, out_dir: Path) -> list[dict]:
    """Each page of a PDF read as a slide is: its text, and its image in *out_dir*.

    The page's title is its first line of text; a PDF has no notes.
    """
    import pymupdf

    out_dir.mkdir(parents=True, exist_ok=True)
    pages = []
    with pymupdf.open(str(pdf_path)) as doc:
        for index, page in enumerate(doc, start=1):
            zoom = IMAGE_WIDTH_PX / page.rect.width
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            name = f"page-{index:03d}.jpg"
            (out_dir / name).write_bytes(pix.tobytes("jpeg", jpg_quality=JPEG_QUALITY))
            text = _clean(page.get_text("text") or "")
            pages.append({"n": index, "title": _page_title(text), "text": text, "notes": "",
                          "links": [], "image": name, "description": ""})
    return pages


def render_pdf_cached(data: bytes, out_dir: Path) -> tuple[Path, dict]:
    """``(out_dir, {"title", "pages"})`` for a PDF's bytes, rendered once.

    *out_dir* is named by the PDF's content hash, so a sync that did not
    change the PDF finds it rendered and does no work.  Written beside and
    renamed into place, so a reader never sees half a rendering.
    """
    import shutil
    import tempfile

    out_dir = Path(out_dir)
    done = out_dir / "pages.json"
    if done.is_file():
        return out_dir, json.loads(done.read_text(encoding="utf-8"))
    out_dir.parent.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix=out_dir.name[:12] + ".", dir=out_dir.parent))
    try:
        source = work / "source.pdf"
        source.write_bytes(data)
        pages = pdf_pages(source, work)
        source.unlink()
        rendered = {"title": pages[0]["title"] if pages else "", "pages": pages}
        (work / "pages.json").write_text(json.dumps(rendered, ensure_ascii=False), encoding="utf-8")
        try:
            work.rename(out_dir)
        except OSError:
            if not done.is_file():      # not a concurrent render that won the race
                raise
    finally:
        if work.exists():
            shutil.rmtree(work, ignore_errors=True)
    return out_dir, json.loads(done.read_text(encoding="utf-8"))


def build_decks(specs: list[DeckSpec], out_root: Path, *, log=print,
                descriptions: dict | None = None) -> list[dict]:
    """Render every paged document under ``out_root/slides/<deck>``; return
    their manifest entries.

    A PDF whose page count differs from its deck's slide count is stale --
    its images would sit beside the wrong slides' text -- so it is skipped,
    loudly, and that deck is served as text alone.

    *descriptions* is the cache ``describe_slides`` writes, keyed by
    ``image_key``; a slide or page whose image is in it gets its description.
    Reading it costs nothing, so every build uses whatever has been described.
    """
    descriptions = descriptions or {}
    entries = []
    for spec in specs:
        deck_dir = out_root / "slides" / spec.id
        deck_dir.mkdir(parents=True)
        warning = ""
        if spec.pptx is None:
            slides = pdf_pages(spec.pdf, deck_dir)
            images = [s["image"] for s in slides]
        else:
            slides = extract_slides(spec.pptx)
            images = []
            if spec.pdf is None:
                warning = "no PDF given; slides are served as text only"
            else:
                pages = pdf_page_count(spec.pdf)
                if pages != len(slides):
                    warning = (f"{spec.pdf.name} has {pages} pages but {spec.pptx.name} has "
                               f"{len(slides)} slides -- re-export the PDF; slides are served "
                               "as text only")
                else:
                    images = render_pages(spec.pdf, deck_dir)
        described = 0
        for slide, image in zip(slides, images or [None] * len(slides)):
            slide["image"] = image
            slide["description"] = ""
            if image:
                entry = descriptions.get(image_key((deck_dir / image).read_bytes()))
                if entry:
                    slide["description"] = entry.get("description", "")
                    described += 1
        source = spec.pptx or spec.pdf
        # The deck's own title slide names it better than its file name does.
        title = spec.title or (slides[0]["title"] if slides and slides[0]["title"] else source.stem)
        (deck_dir / "deck.json").write_text(json.dumps({
            "id": spec.id, "item": spec.item, "unit": spec.unit, "title": title,
            "format": spec.format, "path": spec.path, "source": source.name, "slides": slides,
        }, indent=1, ensure_ascii=False), encoding="utf-8")
        entries.append({
            "id": spec.id, "item": spec.item, "unit": spec.unit, "title": title,
            "format": spec.format, "path": spec.path,
            "slides": len(slides), "images": bool(images), "warning": warning,
        })
        noun = "slides" if spec.format == SLIDES else "pages"
        log(f"  [{spec.id}] {len(slides)} {noun}, "
            f"{'with images' if images else 'text only'}"
            + (f", {described} described" if images else "")
            + (f"  WARNING: {warning}" if warning else ""))
    return entries


def _text_files(part: PartSpec) -> list[tuple[str, Path]]:
    """``(shown path, file)`` for each file of a package part read as text."""
    if part.source_path.is_file():
        suffix = part.source_path.suffix.lower()
        return [] if suffix in (".pptx", ".pdf") else [(part.path, part.source_path)]
    files = []
    for full in sorted(part.source_path.rglob("*")):
        relative = full.relative_to(part.source_path).as_posix()
        if full.is_file() and full.suffix.lower() not in (".pptx", ".pdf"):
            files.append((f"{part.path}/{relative}", full))
    return files


def _build_item(item: ItemSpec, spec: MaterialsSpec, out_root: Path, deck_entries: dict,
                log) -> dict:
    """One package item's manifest entry, its text files copied under
    ``files/<item>/``.  A private item's files are never copied: listing it
    is all the portal can do, so it is all the package carries."""
    from llmgrader.coursemcp.code import read_servable

    policy = spec.policy(item.kind)
    entry = {"id": item.id, "kind": item.kind, "title": item.title,
             "description": item.description, "units": list(item.units),
             "slides": item.slides, "not_slides": item.not_slides}
    parts: list[dict] = []
    if policy == "private":
        entry["parts"] = [{"path": part.path} for part in item.parts]
        log(f"  [{item.id}] {item.kind}, private: listed only, nothing built")
        return entry
    skipped = 0
    for part in item.parts:
        for deck in deck_entries.get(item.id, []):
            if deck["path"] == part.path or deck["path"].startswith(part.path + "/"):
                parts.append({"path": deck["path"], "format": deck["format"], "deck": deck["id"]})
        for shown, full in _text_files(part):
            source = read_servable(full, shown, TEXT, ())
            if source is None:
                if part.source_path.is_file():
                    raise MaterialsError(f"item {item.id!r}: {part.path} cannot be served as text "
                                         "(binary, too large, build output or an unknown type)")
                skipped += 1
                continue
            stored = f"files/{item.id}/{shown}"
            target = out_root / stored
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(full.read_bytes())
            parts.append({"path": shown, "format": TEXT, "file": stored,
                          "language": source.language, "lines": source.last_line})
    if not entry["title"]:
        first_deck = next((d for d in deck_entries.get(item.id, [])), None)
        entry["title"] = first_deck["title"] if first_deck else item.id
    entry["parts"] = parts
    texts = sum(1 for p in parts if p["format"] == TEXT)
    log(f"  [{item.id}] {item.kind}: {len(parts) - texts} paged, {texts} text file(s)"
        + (f", {skipped} not servable and left out" if skipped else ""))
    return entry


def build_materials(spec: MaterialsSpec, package_dir: str | Path, *, log=print,
                    descriptions: dict | None = None, units: list[UnitSpec] | None = None,
                    unit_types: list[dict] | None = None, code: dict | None = None) -> dict:
    """Write the package's material under ``<package_dir>/mcp_materials``;
    return the manifest.

    Package items are built in full.  Mirror items, the ``<items>`` patterns
    and refinements of items only a pattern makes travel as configuration:
    the portal resolves them against its copy of the repo (``catalog.py``).
    *units* (``read_units``) are copied to ``mcp_materials/units/``, each
    with its images folder, and *unit_types* (``read_unit_types``) go into
    the manifest as they are.  So does *code* (``read_code(...).config``),
    the repo the portal keeps a copy of.
    """
    out_root = Path(package_dir) / MATERIALS_DIR
    if out_root.exists():
        import shutil
        shutil.rmtree(out_root)
    out_root.mkdir(parents=True)

    decks = build_decks(spec.decks, out_root, log=log, descriptions=descriptions)
    by_item: dict[str, list[dict]] = {}
    for deck in decks:
        by_item.setdefault(deck["item"], []).append(deck)
    manifest = {
        "version": FORMAT_VERSION,
        "kinds": spec.kinds,
        "material_units": spec.material_units,
        "items": [_build_item(item, spec, out_root, by_item, log)
                  for item in spec.items if item.source == "package"],
        "mirror_items": [
            {"id": item.id, "kind": item.kind, "title": item.title,
             "description": item.description, "units": item.units,
             "slides": item.slides, "not_slides": item.not_slides,
             "parts": [{"path": part.path, "docs": part.docs} for part in item.parts]}
            for item in spec.items if item.source == "mirror"],
        "patterns": spec.patterns,
        "refinements": spec.refinements,
        "decks": decks,
    }
    manifest["units"] = [_build_unit(unit, out_root, log) for unit in units or []]
    manifest["unit_types"] = list(unit_types or [])
    if code:
        # The pointer only: the portal syncs the files itself (read_code).
        manifest["code"] = code
        log(f"  [mirror] {code['repo']} ({code['branch']}): "
            + ", ".join(i["path"] for i in code["includes"]))
    (out_root / MANIFEST).write_text(json.dumps(manifest, indent=1, ensure_ascii=False),
                                     encoding="utf-8")
    return manifest


def _build_unit(spec: UnitSpec, out_root: Path, log) -> dict:
    """Copy one MCP-only unit, and its images, into the package.

    The author writes figures as on the portal, ``/pkg_assets/<stem>_images/``;
    here they live under ``mcp_materials/units/``, so those references are
    rewritten to the package path, and resolve through /pkg_assets exactly
    as a portal unit's do.
    """
    import shutil

    units_dir = out_root / "units"
    units_dir.mkdir(exist_ok=True)
    text = spec.path.read_text(encoding="utf-8")
    images = 0
    if spec.images is not None:
        shutil.copytree(spec.images, units_dir / f"{spec.stem}_images")
        images = sum(1 for p in spec.images.rglob("*") if p.is_file())
        text = text.replace(f"/pkg_assets/{spec.stem}_images/",
                            f"/pkg_assets/{MATERIALS_DIR}/units/{spec.stem}_images/")
    (units_dir / f"{spec.stem}.xml").write_text(text, encoding="utf-8")
    log(f"  [unit] {spec.name}" + (f" ({spec.section})" if spec.section else "")
        + (f", {images} image(s)" if images else ""))
    return {"name": spec.name, "section": spec.section, "file": f"units/{spec.stem}.xml"}


def build_inputs(spec: MaterialsSpec) -> list[Path]:
    """Every package file *spec* builds from: what the package version hashes."""
    return [path for item in spec.items
            if item.source == "package" and spec.policy(item.kind) != "private"
            for part in item.parts for path in (part.source_path, part.pdf) if path]


# ---------------------------------------------------------------------------
# Describe (the one step that costs money)
# ---------------------------------------------------------------------------

DESCRIPTIONS_DIR = "llmgrader_mcp_descriptions"
# The single-file format of the first release; split per deck on first use.
LEGACY_DESCRIPTIONS_FILE = "llmgrader_mcp_descriptions.json"
# Rough per-slide token use, for the estimate --dry-run prints before anything
# is spent: a 1280-px slide image plus its text in, a few sentences out.
EST_TOKENS_IN = 1500
EST_TOKENS_OUT = 200
TEXT_ONLY = "TEXT ONLY"
# Below this much text, a slide is never let off with TEXT ONLY: whatever it
# says is in its figure.  (A figure-only slide was being skipped.)
MIN_TEXT_FOR_TEXT_ONLY = 120

_TEXT_ONLY_RULE = (
    "Only if the slide has no figure, diagram, waveform, table, code, equation "
    "or image of any kind -- nothing but the text below -- reply with exactly: "
    + TEXT_ONLY
)
_DESCRIBE_ALWAYS = (
    "The slide has little text, so its content is in what it shows: always "
    "describe it."
)


def text_only_rule(slide_text: str) -> str:
    return _TEXT_ONLY_RULE if len(slide_text.strip()) >= MIN_TEXT_FOR_TEXT_ONLY else _DESCRIBE_ALWAYS

DESCRIBE_PROMPT = """\
This is {where}, from the course unit "{unit}".
Its own text is below the image.

Describe what the slide shows that its text alone does not: diagrams, block
diagrams, timing waveforms, state diagrams, plots, equations, code and tables
-- what each depicts, its labels and values, and the point it makes.  Use the
technical terms a student would search for (for example "Moore machine",
"handshake", "two's complement").  2 to 5 sentences of plain text, no preamble.

{text_only_rule}

Slide text:
{text}
"""


def descriptions_dir(config_path: str | Path) -> Path:
    """Where the descriptions live: ``<descriptions path>`` in the config, or
    ``llmgrader_mcp_descriptions/`` beside it.  Meant to be committed.

    One JSON file per deck, entries in slide order, so a lecture's
    descriptions can be read top to bottom, a diff shows one lecture, and
    deleting a deck's file redoes just that lecture.
    """
    import xml.etree.ElementTree as ET

    config_path = Path(config_path).resolve()
    try:
        elem = ET.parse(config_path).getroot().find("descriptions")
    except (OSError, ET.ParseError):
        elem = None
    relative = elem.get("path") if elem is not None and elem.get("path") else DESCRIPTIONS_DIR
    return (config_path.parent / relative).resolve()


def _legacy_file(directory: Path) -> Path:
    return Path(directory).parent / LEGACY_DESCRIPTIONS_FILE


def load_descriptions(directory: str | Path) -> dict:
    """Every description, keyed by ``image_key``: the deck files in
    *directory*, over the legacy single file if one is still beside it.

    Matching is by the slide image, not by deck and number, so a renumbered
    slide keeps its description and an edited one is described afresh; an
    identical image in two decks is described once.
    """
    directory = Path(directory)
    cache: dict = {}
    try:
        legacy = json.loads(_legacy_file(directory).read_text(encoding="utf-8"))
        cache.update(legacy.get("slides", {}))
    except (OSError, ValueError):
        pass
    for path in sorted(directory.glob("*.json")) if directory.is_dir() else []:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        for entry in data.get("slides", []):
            cache[entry["image_sha256"]] = {
                "description": entry.get("description", ""), "model": entry.get("model", ""),
                "deck": data.get("deck", path.stem), "slide": entry.get("slide"),
            }
    return cache


def _replace(tmp: Path, path: Path) -> None:
    """``tmp.replace(path)``, waiting out a brief lock.

    On Windows the replace fails while anything else has the file open -- an
    editor, a virus scan, the search indexer, a reader counting entries.
    Those holds are brief, so wait them out rather than fail the slide.
    """
    import time

    for attempt in range(20):
        try:
            tmp.replace(path)
            return
        except PermissionError:
            if attempt == 19:
                raise
            time.sleep(0.25)


def write_deck_descriptions(directory: Path, deck_id: str, slides: list, cache: dict) -> None:
    """Write one deck's file from its *current* slides, in slide order.

    Rewriting from the current slides is what renumbers entries after slides
    are inserted, and drops descriptions of images no longer in the deck.
    *slides* is ``[(number, title, image_key), ...]``.
    """
    entries = [
        {"slide": number, "title": title, "description": cache[key]["description"],
         "model": cache[key].get("model", ""), "image_sha256": key}
        for number, title, key in slides if key in cache
    ]
    if not entries:
        return
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{deck_id}.json"
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps({"version": 2, "deck": deck_id, "slides": entries},
                              indent=1, ensure_ascii=False), encoding="utf-8")
    _replace(tmp, path)


@dataclass
class DescribePlan:
    todo: list            # (deck_meta, slide, image_bytes, key) to describe
    cached: int           # slides with an image already described, left alone
    text_only_decks: list  # deck ids with no usable images, so nothing to describe
    replacing: int = 0    # with force: slides in todo whose description is replaced
    # Every deck with images: deck id -> [(number, title, image_key), ...]
    decks: dict = field(default_factory=dict)


def plan_descriptions(specs: list[DeckSpec], cache: dict, work_dir: Path,
                      *, force: bool = False, redo_empty: bool = False) -> DescribePlan:
    """Render every deck into *work_dir* and work out which slides need a call.

    A slide is sent when its image is not in *cache* -- which is how an edited
    slide gets a new description and an unchanged one does not -- or always,
    with *force*, to redo descriptions (a better model, a poor result).
    *redo_empty* re-sends only the slides whose description is empty -- the
    model judged them text-only -- which is what to rerun after the rule for
    that judgment changes.  Identical images are sent once.
    """
    root = Path(work_dir) / MATERIALS_DIR
    if root.exists():
        import shutil
        shutil.rmtree(root)
    root.mkdir(parents=True)
    metas = build_decks(specs, root, log=lambda *_: None, descriptions=cache)
    todo, cached, text_only, replacing = [], 0, [], 0
    decks: dict = {}
    queued: set[str] = set()
    for meta in metas:
        if not meta["images"]:
            text_only.append(meta["id"])
            continue
        deck_dir = root / "slides" / meta["id"]
        deck = json.loads((deck_dir / "deck.json").read_text(encoding="utf-8"))
        decks[meta["id"]] = []
        for slide in deck["slides"]:
            data = (deck_dir / slide["image"]).read_bytes()
            key = image_key(data)
            decks[meta["id"]].append((slide["n"], slide["title"], key))
            if key in queued:
                continue
            redo = force or (redo_empty and key in cache and not cache[key].get("description"))
            if key in cache and not redo:
                cached += 1
                continue
            replacing += key in cache
            queued.add(key)
            todo.append((meta, slide, data, key))
    return DescribePlan(todo, cached, text_only, replacing, decks)


def estimate_cost(model_spec, slides: int) -> float:
    return slides * (EST_TOKENS_IN * model_spec.usd_per_mtok_in
                     + EST_TOKENS_OUT * model_spec.usd_per_mtok_out) / 1e6


def sync_descriptions(plan: DescribePlan, directory: Path, cache: dict, *, log=print) -> None:
    """Rewrite every deck's file from the current slides, and retire a legacy file.

    Free: no model call.  It renumbers entries after slides moved, and on the
    first run after the per-deck format, splits the old single file -- which is
    then renamed ``.bak`` rather than deleted, since until the deck files are
    committed it is the only copy of what was paid for.
    """
    for deck_id, slides in plan.decks.items():
        write_deck_descriptions(directory, deck_id, slides, cache)
    legacy = _legacy_file(directory)
    if legacy.is_file():
        backup = legacy.with_suffix(".json.bak")
        _replace(legacy, backup)
        log(f"Split {legacy.name} into {Path(directory).name}/<deck>.json; "
            f"the old file is now {backup.name} and can be deleted.")


def describe_slides(plan: DescribePlan, directory: Path, *, model_spec, api_key: str,
                    workers: int = 8, timeout: float = 120, log=print) -> tuple[int, float]:
    """Describe every slide in *plan*; returns ``(described, usd_spent)``.

    A deck's file is rewritten after every slide, so an interruption -- or a
    failed call -- loses nothing already paid for, and re-running picks up
    where it stopped.
    """
    import base64
    from concurrent.futures import ThreadPoolExecutor, as_completed

    from llmgrader.services.answers import make_openai_answer_caller

    cache = load_descriptions(directory)
    sync_descriptions(plan, directory, cache, log=log)
    lock = threading.Lock()
    spent = [0.0]
    done = [0]

    def describe(item):
        meta, slide, data, key = item
        where = (f'slide {slide["n"]} of the lecture deck "{meta["title"]}"'
                 if meta.get("format", SLIDES) == SLIDES
                 else f'page {slide["n"]} of the document "{meta["title"]}"')
        prompt = DESCRIBE_PROMPT.format(
            where=where, unit=meta["unit"],
            text=slide["text"] or "(none)", text_only_rule=text_only_rule(slide["text"]))
        uri = "data:image/jpeg;base64," + base64.b64encode(data).decode("ascii")
        call = make_openai_answer_caller(model=model_spec.id, api_key=api_key,
                                         prompt=prompt, timeout=timeout, images=[uri])
        text, tokens_in, tokens_out = call()
        text = " ".join(text.split())
        description = "" if text.upper().startswith(TEXT_ONLY) else text
        cost = (tokens_in * model_spec.usd_per_mtok_in + tokens_out * model_spec.usd_per_mtok_out) / 1e6
        with lock:
            cache[key] = {"description": description, "model": model_spec.id,
                          "deck": meta["id"], "slide": slide["n"]}
            spent[0] += cost
            done[0] += 1
            # Every deck that shows this image -- usually just the one.
            for deck_id, slides in plan.decks.items():
                if any(k == key for _, _, k in slides):
                    write_deck_descriptions(directory, deck_id, slides, cache)
            if done[0] % 25 == 0 or done[0] == len(plan.todo):
                log(f"  {done[0]}/{len(plan.todo)} described, ${spent[0]:.2f} so far")

    failures = 0
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(describe, item): item for item in plan.todo}
        for future in as_completed(futures):
            try:
                future.result()
            except Exception as exc:  # one bad call must not cost the rest
                failures += 1
                meta, slide, _, _ = futures[future]
                log(f"  failed: {meta['id']} slide {slide['n']}: {exc}")
    if failures:
        log(f"  {failures} slide(s) failed; re-run to retry them -- the rest are saved.")
    return done[0], spent[0]


# ---------------------------------------------------------------------------
# Serve
# ---------------------------------------------------------------------------

_WORD_RE = re.compile(r"[A-Za-z0-9_]+")
_CAMEL_BOUNDARY = re.compile(r"(?<=[a-z0-9])(?=[A-Z])")
# A piece of an identifier shorter than this is not a term: the "s" of
# s_axilite would otherwise match every plural and possessive in the course.
MIN_PART_CHARS = 2
# Hits scoring under this fraction of the best one are dropped: a term the
# best slide shares with a weak one says little about the weak one.
RELATIVE_CUTOFF = 0.3


def _split_word(word: str) -> tuple[str, list[str]]:
    """``(whole, parts)``: a word, lowercased, and the pieces of an identifier.

    snake_case and camelCase split into parts; a plain word ("FIFOs",
    "AXI4") has none, and is a term exactly as before.
    """
    pieces = [p for chunk in word.split("_") for p in _CAMEL_BOUNDARY.split(chunk) if p]
    whole = word.strip("_").lower()
    return whole, ([p.lower() for p in pieces] if len(pieces) > 1 else [])


def tokenize(text: str) -> list[str]:
    """Search terms in *text*.  An identifier counts whole and by its parts:
    ``s_axilite`` is ``s_axilite`` and ``axilite``, so searching for either
    finds it, and the whole identifier -- rare, so heavily weighted -- ranks
    an exact match first."""
    terms = []
    for word in _WORD_RE.findall(text):
        whole, parts = _split_word(word)
        if whole:
            terms.append(whole)
        terms.extend(p for p in parts if len(p) >= MIN_PART_CHARS)
    return terms


def identifiers(text: str) -> set[str]:
    """The identifiers (snake_case, camelCase) in *text*, lowercased."""
    return {whole for whole, parts in map(_split_word, _WORD_RE.findall(text)) if parts}


class TermIndex:
    """BM25 over a set of documents, with identifiers kept whole (tokenize),
    exact identifier matches first, and weak hits cut (RELATIVE_CUTOFF).

    A document is ``(title, [body texts], payload)``; the title counts twice.
    Slides, PDF pages and docs-page sections are ranked in one index, so
    their scores compare (``plans/unified_materials.md``, Tools).
    """

    K1, B = 1.5, 0.75

    def __init__(self, documents: list[tuple[str, list[str], object]]):
        self.docs = []
        for title, bodies, payload in documents:
            terms = tokenize(title) * 2
            for body in bodies:
                terms += tokenize(body or "")
            self.docs.append((Counter(terms), len(terms), payload))
        self.avg_len = (sum(d[1] for d in self.docs) / len(self.docs) or 1) if self.docs else 1

    def search(self, query: str, *, keep=lambda payload: True) -> list[tuple[float, object]]:
        """``(score, payload)`` best first: scored against the whole index, then
        filtered by *keep*, and cut relative to the best hit *keep* allows."""
        terms = tokenize(query)
        wanted_ids = identifiers(query)
        if not self.docs or not terms:
            return []
        df = Counter(t for d in self.docs for t in set(d[0]) if t in terms)
        scored = []
        for counts, length, payload in self.docs:
            if not keep(payload):
                continue
            score = 0.0
            for term in set(terms):
                tf = counts.get(term, 0)
                if not tf:
                    continue
                idf = math.log(1 + (len(self.docs) - df[term] + 0.5) / (df[term] + 0.5))
                score += idf * tf * (self.K1 + 1) / (
                    tf + self.K1 * (1 - self.B + self.B * length / self.avg_len))
            if score > 0:
                # A document with every identifier the query names, whole,
                # ranks above one that only shares their parts.
                exact = all(counts.get(i, 0) for i in wanted_ids)
                scored.append((exact, score, payload))
        if not scored:
            return []
        best = max(score for _, score, _ in scored)
        scored = [item for item in scored if item[1] >= RELATIVE_CUTOFF * best]
        scored.sort(key=lambda item: (not item[0], -item[1]))
        return [(score, payload) for _, score, payload in scored]


def snippet(texts: list[str], query: str, width: int = 160) -> str:
    """A stretch of *texts* around the first query term found in them."""
    text = " ".join(" ".join(t or "" for t in texts).split())
    lower = text.lower()
    hits = [lower.find(t) for t in tokenize(query) if lower.find(t) >= 0]
    start = max(0, min(hits) - 40) if hits else 0
    piece = text[start:start + width]
    return ("…" if start else "") + piece + ("…" if start + width < len(text) else "")


@dataclass
class Materials:
    """One package's MCP material, as its manifest describes it.

    The package items are complete here; mirror items, patterns and
    refinements are configuration that ``catalog.py`` resolves against the
    portal's copy of the repo.
    """

    root: Path
    decks: list[dict] = field(default_factory=list)
    kinds: list[dict] = field(default_factory=list)
    material_units: list[dict] = field(default_factory=list)
    items: list[dict] = field(default_factory=list)           # package items, built
    mirror_items: list[dict] = field(default_factory=list)
    patterns: list[dict] = field(default_factory=list)
    refinements: list[dict] = field(default_factory=list)
    units: list[dict] = field(default_factory=list)           # MCP-only question units
    unit_types: list[dict] = field(default_factory=list)
    code: dict | None = None                                  # <code>: the mirror
    version: int = FORMAT_VERSION
    _deck_cache: dict = field(default_factory=dict)
    _text_cache: dict = field(default_factory=dict)
    _parsed_units: tuple | None = None
    _units_lock: threading.Lock = field(default_factory=threading.Lock)

    def parsed_units(self, *, scratch_dir: str, course_id: str | None) -> tuple[dict, dict]:
        """``(units, unit_metadata)`` for the MCP-only units, parsed once.

        Parsed as the portal parses its own units, so the content tools see
        one shape.  A Materials object lives until the package changes, and
        so does this.
        """
        with self._units_lock:
            if self._parsed_units is None:
                from llmgrader.services.unit_parser import UnitParser

                parser = UnitParser(scratch_dir=scratch_dir, soln_pkg=str(self.root.parent),
                                    course_id=course_id)
                units, metadata, errors = parser.parse_unit_files(
                    [(u["name"], f"{MATERIALS_DIR}/{u['file']}") for u in self.units])
                for error in errors:
                    print(f"[CourseMCP] unit not loaded: {error}")
                self._parsed_units = (units, metadata)
            return self._parsed_units

    def deck_ids(self) -> list[str]:
        return [d["id"] for d in self.decks]

    def deck(self, deck_id: str) -> dict | None:
        if not any(d["id"] == deck_id for d in self.decks):
            return None
        if deck_id not in self._deck_cache:
            path = self.root / "slides" / deck_id / "deck.json"
            self._deck_cache[deck_id] = json.loads(path.read_text(encoding="utf-8"))
        return self._deck_cache[deck_id]

    def image_path(self, deck_id: str, name: str) -> Path | None:
        path = (self.root / "slides" / deck_id / name).resolve()
        slides_root = (self.root / "slides").resolve()
        if not str(path).startswith(str(slides_root) + os.sep) or not path.is_file():
            return None
        return path

    def text_source(self, part: dict):
        """A package text part as a ``code.SourceFile``, read once.  Its file
        comes from the manifest, never from a tool's argument, and is confined
        to ``files/`` anyway."""
        from llmgrader.coursemcp.code import make_source

        stored = part.get("file") or ""
        if stored not in self._text_cache:
            files_root = (self.root / "files").resolve()
            path = (self.root / stored).resolve()
            source = None
            if str(path).startswith(str(files_root) + os.sep) and path.is_file():
                text = path.read_bytes().decode("utf-8-sig", errors="replace")
                if path.suffix == ".ipynb":
                    from llmgrader.coursemcp.code import flatten_notebook
                    text = flatten_notebook(text) or text
                source = make_source(part["path"], TEXT, text)
            self._text_cache[stored] = source
        return self._text_cache[stored]


_CACHE: dict[str, tuple[float, Materials | None]] = {}
_CACHE_LOCK = threading.Lock()


def load_materials(soln_pkg: str | None) -> Materials | None:
    """The package's MCP material, or None if it has none.

    Cached per package directory and reloaded when the manifest changes --
    an upload replaces the package, and with it the manifest.
    """
    if not soln_pkg:
        return None
    root = Path(soln_pkg) / MATERIALS_DIR
    manifest_path = root / MANIFEST
    try:
        mtime = manifest_path.stat().st_mtime
    except OSError:
        return None
    key = str(root.resolve())
    with _CACHE_LOCK:
        cached = _CACHE.get(key)
        if cached and cached[0] == mtime:
            return cached[1]
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        version = manifest.get("version", 1)
        if version < FORMAT_VERSION:
            # No backward compatibility (plans/unified_materials.md): a package
            # built before <materials> serves its exam units, and no material.
            print(f"[CourseMCP] {soln_pkg}: material format {version}, built before "
                  "<materials>; rebuild the package to serve its slides and demos")
        current = version >= FORMAT_VERSION
        materials = Materials(
            root=root, version=version,
            decks=manifest.get("decks", []) if current else [],
            kinds=manifest.get("kinds", []),
            material_units=manifest.get("material_units", []),
            items=manifest.get("items", []),
            mirror_items=manifest.get("mirror_items", []),
            patterns=manifest.get("patterns", []),
            refinements=manifest.get("refinements", []),
            units=manifest.get("units", []),
            unit_types=manifest.get("unit_types", []),
            code=manifest.get("code") if current else None)
        _CACHE[key] = (mtime, materials)
        return materials
